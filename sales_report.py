"""Command-line entry point and Excel writer for sales reports."""

import argparse
import sys

from sales_report_core import (CATEGORY, DATE, PRICE, PRODUCT, QUANTITY, SALES,
                              SalesReportError, load_and_aggregate)


def write_excel(report, output_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as error:
        raise SalesReportError("Excel\u51fa\u529b\u306b\u306fopenpyxl\u304c\u5fc5\u8981\u3067\u3059\u3002python -m pip install openpyxl \u3092\u5b9f\u884c\u3057\u3066\u304f\u3060\u3055\u3044\u3002") from error
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "\u58f2\u4e0a\u30ec\u30dd\u30fc\u30c8"
    sheet.append([PRODUCT, QUANTITY, PRICE, SALES])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in report.rows:
        sheet.append([row["product"], row["quantity"], row["price"], row["sales"]])
    sheet.append([])
    sheet.append(["\u7dcf\u500b\u6570", report.total_quantity])
    sheet.append(["\u7dcf\u58f2\u4e0a", report.total_sales])
    sheet.append(["\u5e73\u5747\u5358\u4fa1", report.average_price])
    _format_sheet(sheet)
    for title, values in (("\u5546\u54c1\u5225", report.product_totals),
                          ("\u65e5\u4ed8\u5225", report.date_totals),
                          ("\u30ab\u30c6\u30b4\u30ea\u5225", report.category_totals)):
        if values:
            summary = workbook.create_sheet(title)
            summary.append([title.replace("\u5225", ""), "\u7dcf\u500b\u6570", "\u7dcf\u58f2\u4e0a"])
            for cell in summary[1]:
                cell.font = Font(bold=True)
            for label, totals in sorted(values.items()):
                summary.append([label, totals["quantity"], totals["sales"]])
            _format_sheet(summary)
    try:
        workbook.save(output_path)
    except (OSError, ValueError) as error:
        raise SalesReportError(f"Excel\u30d5\u30a1\u30a4\u30eb\u3092\u4fdd\u5b58\u3067\u304d\u307e\u305b\u3093: {error}") from error


def _format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else "A1"
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 36)
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.##"


def main(argv=None):
    parser = argparse.ArgumentParser(description="CSV\u304b\u3089\u58f2\u4e0aExcel\u30ec\u30dd\u30fc\u30c8\u3092\u4f5c\u6210\u3057\u307e\u3059")
    parser.add_argument("csv_path", nargs="?", help="\u5165\u529bCSV\u306e\u30d1\u30b9\uff08\u7701\u7565\u6642\u306f\u5bfe\u8a71\u5165\u529b\uff09")
    parser.add_argument("-o", "--output", default="sales_report.xlsx", help="\u51fa\u529bExcel\u306e\u30d1\u30b9")
    args = parser.parse_args(argv)
    csv_path = args.csv_path or input("CSV\u30d5\u30a1\u30a4\u30eb\u306e\u30d1\u30b9\u3092\u5165\u529b\u3057\u3066\u304f\u3060\u3055\u3044\uff1a").strip().strip('"')
    try:
        report = load_and_aggregate(csv_path)
        write_excel(report, args.output)
    except SalesReportError as error:
        print(f"\u30a8\u30e9\u30fc: {error}", file=sys.stderr)
        return 1
    print("\u58f2\u4e0a\u30ec\u30dd\u30fc\u30c8\u3092\u4f5c\u6210\u3057\u307e\u3057\u305f\u3002")
    print(f"\u4fdd\u5b58\u5148\uff1a{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

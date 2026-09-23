import csv
import tempfile
import unittest
from pathlib import Path

from sales_report_core import (CATEGORY, DATE, PRICE, PRODUCT, QUANTITY, SALES,
                              SalesReportError, load_and_aggregate)
from sales_report import write_excel


class SalesReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "input.csv"

    def write_csv(self, headers, rows):
        with self.path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)

    def test_normal_csv_and_totals(self):
        self.write_csv([PRODUCT, QUANTITY, PRICE], [["apple", 3, 100], ["banana", 2, 150]])
        report = load_and_aggregate(self.path)
        self.assertEqual((report.total_quantity, report.total_sales, report.average_price), (5, 600, 120))
        self.assertEqual(report.product_totals["apple"], {"quantity": 3, "sales": 300})

    def test_empty_rows_return_zero_report(self):
        self.write_csv([PRODUCT, QUANTITY, PRICE], [])
        report = load_and_aggregate(self.path)
        self.assertEqual((report.rows, report.total_quantity, report.total_sales), ([], 0, 0))

    def test_completely_empty_csv_returns_zero_report(self):
        self.path.write_text("", encoding="utf-8")
        report = load_and_aggregate(self.path)
        self.assertEqual((report.rows, report.total_quantity, report.total_sales), ([], 0, 0))

    def test_missing_required_column(self):
        self.write_csv([PRODUCT, QUANTITY], [["apple", 3]])
        with self.assertRaisesRegex(SalesReportError, PRICE):
            load_and_aggregate(self.path)

    def test_invalid_number_reports_line(self):
        self.write_csv([PRODUCT, QUANTITY, PRICE], [["apple", "bad", 100]])
        with self.assertRaises(SalesReportError):
            load_and_aggregate(self.path)

    def test_optional_date_and_category_aggregates(self):
        self.write_csv([PRODUCT, QUANTITY, PRICE, DATE, CATEGORY], [["apple", 2, 100, "2026-09-01", "fruit"]])
        report = load_and_aggregate(self.path)
        self.assertEqual(report.date_totals["2026-09-01"]["sales"], 200)
        self.assertEqual(report.category_totals["fruit"]["quantity"], 2)

    def test_excel_output_contains_japanese_sheet_and_headers(self):
        self.write_csv([PRODUCT, QUANTITY, PRICE], [["\\u308a\\u3093\\u3054", 1, 100]])
        output = Path(self.temp.name) / "report.xlsx"
        write_excel(load_and_aggregate(self.path), output)
        from openpyxl import load_workbook
        sheet = load_workbook(output).active
        self.assertEqual(sheet.title, "\u58f2\u4e0a\u30ec\u30dd\u30fc\u30c8")
        self.assertEqual(sheet.cell(1, 1).value, PRODUCT)
        self.assertEqual(sheet.cell(2, 1).value, "\\u308a\\u3093\\u3054")


if __name__ == "__main__":
    unittest.main()
